"""Shared helpers for the Excel workbook generators (openpyxl)."""
import os, sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as col
from openpyxl.worksheet.datavalidation import DataValidation

HERE = os.path.dirname(os.path.abspath(__file__))
MODELS = os.path.join(HERE, "..", "models")
sys.path.insert(0, MODELS)
sys.path.insert(0, os.path.join(HERE, "..", "physics"))

FONT = "Arial"
F = Font(name=FONT, size=10)
F_IN = Font(name=FONT, size=10, color="0000FF")
F_LINK = Font(name=FONT, size=10, color="008000")
F_B = Font(name=FONT, size=10, bold=True)
F_H = Font(name=FONT, size=10, bold=True, color="FFFFFF")
F_T = Font(name=FONT, size=14, bold=True)
F_NOTE = Font(name=FONT, size=9, italic=True, color="555555")
FILL_IN = PatternFill("solid", fgColor="FFF6CC")
FILL_H = PatternFill("solid", fgColor="2F3B40")
FILL_OUT = PatternFill("solid", fgColor="E8F1F6")
THIN = Side(style="thin", color="C8CED0")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def new_wb():
    wb = Workbook()
    wb.remove(wb.active)
    return wb


def sheet(wb, name, widths=None):
    ws = wb.create_sheet(name)
    ws.sheet_view.showGridLines = False
    for c, w in (widths or {}).items():
        ws.column_dimensions[c].width = w
    return ws


def put(ws, ref, value, font=F, fill=None, fmt=None, bold=False, align=None, border=False):
    c = ws[ref]
    c.value = value
    c.font = F_B if bold and font is F else font
    if fill:
        c.fill = fill
    if fmt:
        c.number_format = fmt
    if align:
        c.alignment = Alignment(horizontal=None if align == "wrap" else align, vertical="center", wrap_text=align == "wrap")
    if border:
        c.border = BOX
    return c


def inp(ws, ref, value, fmt=None):
    return put(ws, ref, value, F_IN, FILL_IN, fmt, border=True)


def out(ws, ref, formula, fmt=None, bold=False):
    return put(ws, ref, formula, F_B if bold else F, FILL_OUT, fmt, border=True)


def header(ws, row, col0, labels, widths=None):
    for i, lab in enumerate(labels):
        c = ws.cell(row=row, column=col0 + i, value=lab)
        c.font = F_H; c.fill = FILL_H
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 30


def note(ws, ref, text):
    c = ws[ref]; c.value = text; c.font = F_NOTE
    c.alignment = Alignment(wrap_text=False, vertical="top")


def readme(wb, title, lines, legend=True):
    ws = sheet(wb, "Read me", {"A": 3, "B": 110})
    put(ws, "B2", title, F_T)
    r = 4
    for ln in lines:
        c = ws.cell(row=r, column=2, value=ln); c.font = F_B if ln.isupper() or ln.endswith(":") else F
        c.alignment = Alignment(wrap_text=True, vertical="top")
        r += 1
    if legend:
        r += 1
        ws.cell(row=r, column=2, value="Colour key:").font = F_B; r += 1
        c = ws.cell(row=r, column=2, value="Blue text on yellow: an input you can change."); c.font = F_IN; c.fill = FILL_IN; r += 1
        c = ws.cell(row=r, column=2, value="Black text on pale blue: a calculated result."); c.font = F; c.fill = FILL_OUT; r += 1
        c = ws.cell(row=r, column=2, value="Plain black: calculation steps and data."); c.font = F; r += 1
    return ws


def dv_list(ws, ref, formula):
    dv = DataValidation(type="list", formula1=formula, allow_blank=False)
    ws.add_data_validation(dv); dv.add(ref)
    return dv
