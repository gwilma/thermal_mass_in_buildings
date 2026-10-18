"""Workbook 2: Conduction Lab (implicit finite differences and the five-node
lump, stepped in cell formulas).

The element is meshed by formulas from the Library sheet. Each method builds
its system matrix A = C/dt + G, inverts it once with MINVERSE, and steps
T(t+dt) = A^-1 (C/dt T(t) + sources) every 5 minutes for four days,
exactly as models/elements.py step_response does.
"""
import os
import numpy as np
from openpyxl.chart import LineChart, Reference
from openpyxl.worksheet.formula import ArrayFormula

from xlcommon import *  # noqa: F401,F403
from build_buildup import library_sheet, materials_sheet, classes_sheet, mat_col, LIB_FIRST, LIB_LAST, MAT_FIRST, MAT_LAST, MATR, N_SLOTS
from buildups import LIB
from elements import step_response

NF = 60          # fine-mesh node slots
N5 = 25          # five-node slots (5 + 20 soil nodes for ground floors)
DT = 300
STEPS = 4 * 288  # four days
LIBK = f"Library!$A${LIB_FIRST}:$A${LIB_LAST}"


def setup_sheet(wb):
    ws = sheet(wb, "Setup", {"A": 30, "B": 14, "C": 60})
    put(ws, "A1", "Conduction Lab: test set-up", F_T)
    note(ws, "A2", "Change the blue cells. The four-day run recalculates (allow a few seconds).")
    rows = [("Element key", "W2", "Any key on the Library sheet, or Custom"),
            ("Mesh cell size, mm", 10, "10 or 20 (the sheet has room for 60 nodes)"),
            ("Five-node class", "rule", "rule, I, E, IE, D or M"),
            ("Room mean, °C", 24, ""), ("Room amplitude, K", 4, ""), ("Room peak hour", 16, ""),
            ("Far-side temperature, °C", 15, "15 for outdoor-facing elements, 24 for internal ones in the Python runs"),
            ("Time step, s", DT, "fixed: the run sheets have 1,152 rows")]
    for i, (lab, v, n) in enumerate(rows):
        r = 4 + i
        put(ws, f"A{r}", lab, bold=True)
        (put if i == 7 else inp)(ws, f"B{r}", v)
        note(ws, f"C{r}", n)
    dv_list(ws, "B4", f"={LIBK}")
    dv_list(ws, "B5", '"10,20"')
    dv_list(ws, "B6", '"rule,I,E,IE,D,M"')
    put(ws, "A13", "Derived", bold=True)
    der = [("Row in library", f"=MATCH(B4,{LIBK},0)"),
           ("Build-up", f"=INDEX(Library!$B${LIB_FIRST}:$B${LIB_LAST},B14)"),
           ("Far side", f"=INDEX(Library!$D${LIB_FIRST}:$D${LIB_LAST},B14)"),
           ("Rsi", f"=INDEX(Library!$E${LIB_FIRST}:$E${LIB_LAST},B14)"),
           ("Rse", '=IF(B16="ext",0.04,0)'),
           ("Class used", f'=IF(B6="rule",INDEX(Library!$F${LIB_FIRST}:$F${LIB_LAST},B14),B6)'),
           ("Far-side conductance W/m²K", '=IF(B16="ext",IF(B18>0,1/B18,1000000),IF(B16="ground",1000000,0))'),
           ("Room-coupled far face (1/0)", '=IF(B16="room",1,0)'),
           ("Far-side source W/m²", "=B20*B10+B21*B10/B17")]
    for i, (lab, f) in enumerate(der):
        put(ws, f"A{14 + i}", lab); out(ws, f"B{14 + i}", f)
    ws["B15"].alignment = ws["B15"].alignment.copy(horizontal="left")
    return ws


def layers_block(ws):
    """Layer table on the Mesh sheet: rows 5..13 (8 slots + soil)."""
    header(ws, 4, 1, ["Layer", "Material", "Type", "d m", "λ", "ρ", "c", "R", "cells", "h m", "g cell W/m²K", "C cell J/m²K", "first cell − 1", "R part"])
    dx = "(Setup!$B$5/1000)"
    for i in range(1, N_SLOTS + 2):
        r = 4 + i
        if i <= N_SLOTS:
            a, b = mat_col(i)
            put(ws, f"A{r}", f"Layer {i}")
            put(ws, f"B{r}", f'=IF(INDEX(Library!${a}${LIB_FIRST}:${a}${LIB_LAST},Setup!$B$14)="","",INDEX(Library!${a}${LIB_FIRST}:${a}${LIB_LAST},Setup!$B$14))')
            val = f"INDEX(Library!${b}${LIB_FIRST}:${b}${LIB_LAST},Setup!$B$14)"
            put(ws, f"C{r}", f'=IF(B{r}="","none",IF(B{r}="Cavity","res","mass"))')
            put(ws, f"D{r}", f'=IF(C{r}="mass",{val}/1000,0)')
            put(ws, f"H{r}", f'=IF(C{r}="res",{val},0)')
            put(ws, f"I{r}", f'=IF(C{r}="mass",MAX(1,CEILING(D{r}/{dx}-1E-9,1)),IF(C{r}="res",1,0))')
        else:
            put(ws, f"A{r}", "Soil (ground floors)")
            put(ws, f"B{r}", '=IF(Setup!$B$16="ground","clay soil","")')
            put(ws, f"C{r}", f'=IF(B{r}="","none","mass")')
            put(ws, f"D{r}", f'=IF(C{r}="mass",1,0)'); put(ws, f"H{r}", 0)
            put(ws, f"I{r}", f'=IF(C{r}="mass",20,0)')
        for cc, k in (("E", "B"), ("F", "C"), ("G", "D")):
            put(ws, f"{cc}{r}", f'=IF(C{r}="mass",INDEX(Materials!${k}${MAT_FIRST}:${k}${MAT_LAST},MATCH(B{r},{MATR},0)),0)')
        put(ws, f"J{r}", f'=IF(I{r}>0,D{r}/I{r},0)')
        put(ws, f"K{r}", f'=IF(C{r}="mass",E{r}/J{r},IF(C{r}="res",1/H{r},0))')
        put(ws, f"L{r}", f'=IF(C{r}="mass",F{r}*G{r}*J{r},0)')
        put(ws, f"M{r}", "=0" if i == 1 else f"=M{r - 1}+I{r - 1}")
        put(ws, f"N{r}", f'=IF(C{r}="mass",D{r}/E{r},H{r})')
    put(ws, "A15", "Cells in total", bold=True); out(ws, "B15", "=SUM(I5:I13)")
    put(ws, "A16", "Nodes", bold=True); out(ws, "B16", "=B15+1")
    put(ws, "C16", '=IF(B16>60,"Too many nodes: use 20 mm cells","")', F_B)


def fd_mesh_sheet(wb):
    ws = sheet(wb, "FD mesh", {"A": 20, "B": 14})
    put(ws, "A1", "Fine mesh: layers → cells → nodes", F_T)
    layers_block(ws)
    # cells j = 1..NF-1 in rows 20.., nodes k = 0..NF-1 in rows 20..
    header(ws, 19, 1, ["Cell j", "layer", "g W/m²K", "C J/m²K", "", "Node k", "valid", "C node J/m²K", "g to next", "C/dt", "diagonal of A"])
    for j in range(1, NF):
        r = 19 + j
        put(ws, f"A{r}", j)
        put(ws, f"B{r}", f'=IF(A{r}<=$B$15,COUNTIF($M$5:$M$13,"<="&(A{r}-1)),0)')
        put(ws, f"C{r}", f"=IF(B{r}>0,INDEX($K$5:$K$13,B{r}),0)")
        put(ws, f"D{r}", f"=IF(B{r}>0,INDEX($L$5:$L$13,B{r}),0)")
    rsi = "Setup!$B$17"
    for k in range(NF):
        r = 20 + k
        put(ws, f"F{r}", k)
        put(ws, f"G{r}", f"=IF(F{r}<=$B$15,1,0)")
        capL = f"D{r - 1}" if k >= 1 else "0"
        capR = f"D{r}" if k < NF - 1 else "0"
        put(ws, f"H{r}", f"=IF(G{r}=1,({capL}+{capR})/2,0)")
        put(ws, f"I{r}", f"=IF(F{r}<$B$15,C{r},0)" if k < NF - 1 else "=0")
        put(ws, f"J{r}", f"=H{r}/Setup!$B$11")
        gL = f"I{r - 1}" if k >= 1 else "0"
        bc = f"+IF(F{r}=0,1/{rsi},0)+IF(F{r}=$B$15,Setup!$B$20+Setup!$B$21/{rsi},0)"
        put(ws, f"K{r}", f"=IF(G{r}=1,J{r}+{gL}+I{r}{bc},1)")
    # horizontal copy of C/dt for SUMPRODUCT
    put(ws, "A84", "C/dt as a row", bold=True)
    for k in range(NF):
        put(ws, f"{col(2 + k)}84", f"=J{20 + k}")
    return ws


def five_mesh_sheet(wb):
    ws = sheet(wb, "5N mesh", {"A": 26, "B": 14})
    put(ws, "A1", "Five-node lump: nodes", F_T)
    note(ws, "A2", "Node 0 is the room-side node (class node 5). Ground floors add the same 1 m of soil as the fine mesh, in 20 cells of 50 mm.")
    put(ws, "A4", "Rc m²K/W", bold=True); out(ws, "B4", "=SUM('FD mesh'!N5:N12)")
    put(ws, "A5", "κ total J/m²K", bold=True); out(ws, "B5", "=SUMPRODUCT('FD mesh'!F5:F12,'FD mesh'!G5:G12,'FD mesh'!D5:D12)")
    put(ws, "A6", "Ground floor (1/0)", bold=True); out(ws, "B6", '=IF(Setup!$B$16="ground",1,0)')
    put(ws, "A7", "Last node", bold=True); out(ws, "B7", "=IF(B6=1,24,4)")
    soilC = 1500 * 2000 * 0.05; soilG = 1.5 / 0.05
    header(ws, 9, 1, ["Node k", "valid", "C node J/m²K", "g to next", "C/dt", "diagonal of A"])
    cls = "Setup!$B$19"
    gl = ["6/$B$4", "3/$B$4", "3/$B$4", "6/$B$4"]
    for k in range(N5):
        r = 10 + k
        put(ws, f"A{r}", k)
        put(ws, f"B{r}", f"=IF(A{r}<=$B$7,1,0)")
        if k < 5:
            share = f"INDEX(Classes!$B$5:$F$9,MATCH({cls},Classes!$A$5:$A$9,0),{5 - k})"
            put(ws, f"C{r}", f"={share}*$B$5" + (f"+$B$6*{soilC / 2}" if k == 4 else ""))
        else:
            put(ws, f"C{r}", f"=IF(B{r}=1,{soilC / 2 if k == N5 - 1 else soilC},0)")
        if k < 4:
            put(ws, f"D{r}", f"={gl[k]}")
        elif k < N5 - 1:
            put(ws, f"D{r}", f"=IF(A{r}<$B$7,{soilG},0)")
        else:
            put(ws, "D" + str(r), "=0")
        put(ws, f"E{r}", f"=C{r}/Setup!$B$11")
        gL = f"D{r - 1}" if k >= 1 else "0"
        bc = f"+IF(A{r}=0,1/Setup!$B$17,0)+IF(A{r}=$B$7,Setup!$B$20+Setup!$B$21/Setup!$B$17,0)"
        put(ws, f"F{r}", f"=IF(B{r}=1,E{r}+{gL}+D{r}{bc},1)")
    put(ws, "A37", "C/dt as a row", bold=True)
    for k in range(N5):
        put(ws, f"{col(2 + k)}37", f"=E{10 + k}")
    return ws


def matrix_sheets(wb, name, n, mesh, r_node0, diag_col, g_col, valid_col):
    A = sheet(wb, f"{name} A")
    put(A, "A1", f"{name}: system matrix A = C/dt + G (with boundary conductances)", F_B)
    for i in range(n):
        for j in range(n):
            ref = f"{col(1 + j)}{3 + i}"
            ri = r_node0 + i
            if i == j:
                A[ref] = f"='{mesh}'!{diag_col}{ri}"
            elif j == i + 1:
                A[ref] = f"=IF('{mesh}'!{valid_col}{r_node0 + j}=1,-'{mesh}'!{g_col}{ri},0)"
            elif j == i - 1:
                A[ref] = f"=IF('{mesh}'!{valid_col}{ri}=1,-'{mesh}'!{g_col}{r_node0 + j},0)"
            else:
                A[ref] = 0
    I = sheet(wb, f"{name} inverse")
    put(I, "A1", f"{name}: MINVERSE of the system matrix (computed once)", F_B)
    rng = f"A3:{col(n)}{2 + n}"
    I["A3"] = ArrayFormula(rng, f"=MINVERSE('{name} A'!A3:{col(n)}{2 + n})")
    return A, I


def run_sheet(wb, name, n, mesh, last_ref, cdt_row):
    ws = sheet(wb, f"{name} run", {"A": 8, "B": 9, "C": 9, "D": 11})
    put(ws, "A1", f"{name}: four days at 5-minute steps (implicit Euler)", F_B)
    header(ws, 3, 1, ["step", "hour", "T room °C", "q into surface W/m²"] + [f"T{k}" for k in range(n)])
    c0 = 5  # column of node 0
    first = 4
    # initial state (row 4): all at far-side temperature
    put(ws, "A4", "start")
    for k in range(n):
        put(ws, f"{col(c0 + k)}4", "=Setup!$B$10")
    inv = f"'{name} inverse'"
    lastc = col(n)
    for s in range(STEPS):
        r = first + 1 + s
        ws[f"A{r}"] = s
        ws[f"B{r}"] = f"=A{r}*Setup!$B$11/3600"
        ws[f"C{r}"] = f"=Setup!$B$7+Setup!$B$8*COS(2*PI()*(B{r}-Setup!$B$9)/24)"
        ws[f"D{r}"] = f"=(C{r}-{col(c0)}{r})/Setup!$B$17"
        prev = f"${col(c0)}{r - 1}:${col(c0 + n - 1)}{r - 1}"
        for k in range(n):
            irow = f"{inv}!$A${3 + k}:${lastc}${3 + k}"
            ws[f"{col(c0 + k)}{r}"] = (f"=SUMPRODUCT({irow},'{mesh}'!$B${cdt_row}:${col(1 + n)}${cdt_row},{prev})"
                                      f"+{inv}!$A${3 + k}*$C{r}/Setup!$B$17"
                                      f"+INDEX({irow},1,{last_ref}+1)*Setup!$B$22")
    return ws


def results_sheet(wb):
    ws = sheet(wb, "Results", {"A": 34, "B": 14, "C": 14, "D": 14, "E": 4, "F": 8, "G": 46, "H": 13, "I": 13, "J": 13, "K": 13})
    put(ws, "A1", "Conduction Lab: results", F_T)
    put(ws, "A2", "=Setup!B4&\": \"&Setup!B15", F_B)
    header(ws, 4, 1, ["", "Fine mesh", "Five-node", "Five-node ÷ mesh"])
    last0, last1 = 5 + STEPS - 288, 4 + STEPS  # rows of the last day on the run sheets
    rng = lambda sh, c: f"'{sh} run'!{c}{last0}:{c}{last1}"
    hour_ok = lambda sh: f"(MOD({rng(sh, 'B')},24)>=9)*(MOD({rng(sh, 'B')},24)<18)"
    put(ws, "A5", "Absorbed 09:00–18:00, last day, Wh/m²")
    out(ws, "B5", f"=SUMPRODUCT({hour_ok('FD')},{rng('FD', 'D')})*Setup!$B$11/3600", "0", True)
    out(ws, "C5", f"=SUMPRODUCT({hour_ok('5N')},{rng('5N', 'D')})*Setup!$B$11/3600", "0", True)
    out(ws, "D5", "=C5/B5", "0.00")
    put(ws, "A6", "Peak flux into surface, W/m²")
    out(ws, "B6", f"=MAX({rng('FD', 'D')})", "0.0"); out(ws, "C6", f"=MAX({rng('5N', 'D')})", "0.0"); out(ws, "D6", "=C6/B6", "0.00")
    put(ws, "A7", "Hour of peak flux")
    out(ws, "B7", f"=MOD(INDEX({rng('FD', 'B')},MATCH(B6,{rng('FD', 'D')},0)),24)", "0.0")
    out(ws, "C7", f"=MOD(INDEX({rng('5N', 'B')},MATCH(C6,{rng('5N', 'D')},0)),24)", "0.0")
    put(ws, "A8", "Nodes used")
    out(ws, "B8", "='FD mesh'!B16"); out(ws, "C8", "='5N mesh'!B7+1")
    put(ws, "A10", "Last-day flux (for the chart)", bold=True)
    header(ws, 11, 1, ["hour", "fine mesh W/m²", "five-node W/m²", "room °C"])
    for i in range(288):
        r = 12 + i
        ws[f"A{r}"] = f"=MOD('FD run'!B{last0 + i},24)"
        ws[f"B{r}"] = f"='FD run'!D{last0 + i}"
        ws[f"C{r}"] = f"='5N run'!D{last0 + i}"
        ws[f"D{r}"] = f"='FD run'!C{last0 + i}"
    ch = LineChart(); ch.title = "Heat flux into the surface, last day"; ch.y_axis.title = "W/m²"; ch.x_axis.title = "hour"
    ch.add_data(Reference(ws, min_col=2, max_col=3, min_row=11, max_row=11 + 288), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=1, min_row=12, max_row=11 + 288))
    ch.x_axis.tickLblSkip = 36; ch.height = 9; ch.width = 18
    ws.add_chart(ch, "F4")
    # Python reference for every build-up at the default set-up
    r0 = 25
    put(ws, f"G{r0 - 1}", "Python reference (default set-up, 10 mm cells, library class)", F_B)
    header(ws, r0, 7, ["Build-up", "Absorbed fine Wh/m²", "Absorbed 5-node Wh/m²", "Peak fine W/m²", "Peak 5-node W/m²"])
    t = np.arange(0, 4 * 86400, DT); Tr = 24 + 4 * np.cos(2 * np.pi * (t / 3600 - 16) / 24)
    h = (t / 3600) % 24; last = t >= t[-1] - 86400 + DT; sel = last & (h >= 9) & (h < 18)
    for i, (k, e) in enumerate(LIB.items()):
        Tb = 15.0 if e.back == "ext" else 24.0
        qf = step_response(e, Tr, DT, "fd", 0.01, Tback=Tb); q5 = step_response(e, Tr, DT, "5node", Tback=Tb)
        r = r0 + 1 + i
        put(ws, f"G{r}", f"{k} · {e.name}")
        for j, v in enumerate([qf[sel].sum() * DT / 3600, q5[sel].sum() * DT / 3600, qf[last].max(), q5[last].max()]):
            put(ws, f"{col(8 + j)}{r}", round(float(v), 2), fmt="0.0")
    note(ws, f"G{r0 + len(LIB) + 2}", "Set the far-side temperature to 24 °C for internal elements (partitions, floors) to match these rows.")
    return ws


def build(path):
    wb = new_wb()
    readme(wb, "Conduction Lab (thermal mass toolkit, model 2)", [
        "What it does:",
        "Drives one element with a sinusoidal room temperature for four days and steps the heat flow through it every five minutes, two ways: a fine finite-difference mesh (the reference) and the ISO 52016-1 / HEM five-node lump. It reports the heat taken in between 09:00 and 18:00 on the last day and the peak flux.",
        "",
        "How to use it:",
        "1. Setup: choose the element (any Library key, or Custom), mesh cell size, five-node class, the room temperature cycle and the far-side temperature.",
        "2. Results: absorbed heat, peak flux and timing for both methods, a chart of the last day, and the Python model's values for every build-up at the default set-up.",
        "",
        "How it works:",
        "FD mesh and 5N mesh build node capacities and conductances from the Library layers with formulas. 'FD A' and '5N A' hold the system matrix C/dt + G; the inverse sheets invert it once with MINVERSE; the run sheets step T(t+dt) = A⁻¹ (C/dt · T(t) + sources) with SUMPRODUCT.",
        "This is the same implicit Euler scheme as models/elements.py step_response, so with the same cell size the numbers match the Python model.",
        "The fine mesh has room for 60 nodes: use 10 mm cells for most elements and 20 mm for thick ones. A warning appears on FD mesh if the element needs more.",
        "",
        "Recalculation: about 100,000 formulas. Excel recalculates in a few seconds; if calculation is set to manual, press F9.",
    ])
    setup_sheet(wb)
    results_sheet(wb)
    fd_mesh_sheet(wb)
    five_mesh_sheet(wb)
    matrix_sheets(wb, "FD", NF, "FD mesh", 20, "K", "I", "G")
    matrix_sheets(wb, "5N", N5, "5N mesh", 10, "F", "D", "B")
    run_sheet(wb, "FD", NF, "FD mesh", "'FD mesh'!$B$15", 84)
    run_sheet(wb, "5N", N5, "5N mesh", "'5N mesh'!$B$7", 37)
    library_sheet(wb); materials_sheet(wb); classes_sheet(wb)
    wb.save(path)


if __name__ == "__main__":
    os.makedirs(os.path.join(HERE, "out"), exist_ok=True)
    build(os.path.join(HERE, "out", "conduction-lab.xlsx"))
    print("written")
