"""Workbook 1: Build-up Explorer (ISO 13786 and five-node, in cell formulas).

Every result is a live formula: the ISO 13786 heat-transfer matrices are
built with Excel's complex-number functions (COMPLEX, IMPRODUCT, IMSUM,
IMDIV, IMABS, IMARGUMENT) and multiplied layer by layer; the five-node lump
is the same chain made of resistances and lumped capacitances.
"""
import os
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment

from xlcommon import *  # noqa: F401,F403
from buildups import LIB, MAT
import csv

N_SLOTS = 8
LIB_FIRST, LIB_LAST = 5, 5 + len(LIB)  # last row is "Custom"
MAT_FIRST = 5
MATS = [m for m in MAT]
MAT_LAST = MAT_FIRST + len(MATS) - 1
CLASSES = {"I": [0, 0, 0, 0, 1], "E": [1, 0, 0, 0, 0], "IE": [.5, 0, 0, 0, .5], "D": [1/8, 1/4, 1/4, 1/4, 1/8], "M": [0, 0, 1, 0, 0]}
MATR = f"Materials!$A${MAT_FIRST}:$A${MAT_LAST}"


def mat_col(i):  # Library column letters for slot i (1-based)
    return col(7 + 2 * (i - 1)), col(8 + 2 * (i - 1))


def materials_sheet(wb):
    ws = sheet(wb, "Materials", {"A": 24, "B": 12, "C": 12, "D": 12, "E": 14, "F": 16, "G": 16})
    put(ws, "A1", "Materials", F_T)
    note(ws, "A2", "Typical design values (CIBSE Guide A, ISO 10456, manufacturers' data). Edit the blue cells to use product data.")
    header(ws, 4, 1, ["Material", "λ W/mK", "ρ kg/m³", "c J/kgK", "α mm²/s", "b J/m²K s½", "δ (24 h) mm"])
    for i, m in enumerate(MATS):
        r = MAT_FIRST + i
        lam, rho, c = MAT[m]
        put(ws, f"A{r}", m)
        inp(ws, f"B{r}", lam); inp(ws, f"C{r}", rho); inp(ws, f"D{r}", c)
        put(ws, f"E{r}", f"=B{r}/(C{r}*D{r})*1E6", fmt="0.00")
        put(ws, f"F{r}", f"=SQRT(B{r}*C{r}*D{r})", fmt="0")
        put(ws, f"G{r}", f"=SQRT(B{r}*86400/(PI()*C{r}*D{r}))*1000", fmt="0")
    put(ws, "I4", "Layer choices", F_B)
    for i, m in enumerate(["Cavity"] + MATS):
        put(ws, f"I{MAT_FIRST + i}", m)
    ws.column_dimensions["I"].width = 22
    note(ws, f"A{MAT_LAST + 2}", "α = λ/ρc (diffusivity); b = √(λρc) (effusivity); δ = √(λT/πρc), the periodic penetration depth for a 24 h cycle.")


def library_sheet(wb):
    w = {"A": 7, "B": 52, "C": 10, "D": 10, "E": 7, "F": 7}
    for i in range(1, N_SLOTS + 1):
        a, b = mat_col(i); w[a] = 20; w[b] = 8
    ws = sheet(wb, "Library", w)
    put(ws, "A1", "Build-up library", F_T)
    note(ws, "A2", "Layers run from the room side outwards. Value is thickness in mm for a material, or resistance in m²K/W for a 'Cavity'. Far side: ext, ground, adiabatic or room. The last row is yours to edit.")
    labels = ["Key", "Build-up", "Kind", "Far side", "Rsi", "Class (rule)"]
    for i in range(1, N_SLOTS + 1):
        labels += [f"Layer {i}", "mm or R"]
    header(ws, 4, 1, labels)
    rows = list(LIB.values())
    for k, e in enumerate(rows + [None]):
        r = LIB_FIRST + k
        if e is None:  # Custom row, seeded with W1
            e0 = LIB["W1"]
            vals = ["Custom", "Your build-up (edit this row)", "wall", "ext", 0.13, e0.cls]
            layers = e0.layers
            for j, v in enumerate(vals):
                inp(ws, f"{col(j + 1)}{r}", v)
        else:
            vals = [e.key, e.name, e.kind, e.back, e.rsi, e.cls]
            layers = e.layers
            for j, v in enumerate(vals):
                put(ws, f"{col(j + 1)}{r}", v)
        for i, lay in enumerate(layers, start=1):
            a, b = mat_col(i)
            if len(lay) == 2:
                (inp if e is None else put)(ws, f"{a}{r}", "Cavity"); (inp if e is None else put)(ws, f"{b}{r}", lay[1])
            else:
                (inp if e is None else put)(ws, f"{a}{r}", lay[0]); (inp if e is None else put)(ws, f"{b}{r}", round(lay[1] * 1000, 2))
        if e is None:
            for i in range(len(layers) + 1, N_SLOTS + 1):
                a, b = mat_col(i); inp(ws, f"{a}{r}", None); inp(ws, f"{b}{r}", None)
    # material choices for the custom row
    for i in range(1, N_SLOTS + 1):
        dv_list(ws, f"{mat_col(i)[0]}{LIB_LAST}", f"=Materials!$I${MAT_FIRST}:$I${MAT_FIRST + len(MATS)}")
    ws.freeze_panes = "C5"
    return ws


def chain_block(ws, r0, key_ref, period_ref, title):
    """ISO 13786 chain for the build-up named in key_ref. Returns refs dict.
    Columns: A label, B type, C material, D d m, E λ, F ρ, G c, H R, I δ, J x,
    K..N matrix (z11 z12 z21 z22), O..R running product, S resistance, T ρcd."""
    put(ws, f"A{r0}", title, F_B)
    idx = f"$U${r0}"
    put(ws, f"T{r0}", "row →", F_NOTE)
    put(ws, idx.replace("$", ""), f"=MATCH({key_ref},Library!$A${LIB_FIRST}:$A${LIB_LAST},0)")
    back = f"$W${r0}"; rsi = f"$Y${r0}"
    put(ws, f"V{r0}", "far side", F_NOTE); put(ws, back.replace("$", ""), f"=INDEX(Library!$D${LIB_FIRST}:$D${LIB_LAST},{idx})")
    put(ws, f"X{r0}", "Rsi", F_NOTE); put(ws, rsi.replace("$", ""), f"=INDEX(Library!$E${LIB_FIRST}:$E${LIB_LAST},{idx})")
    Tsec = f"({period_ref}*3600)"
    labels = ["", "Type", "Material", "d m", "λ", "ρ", "c", "R m²K/W", "δ m", "x = d/δ", "z11", "z12", "z21", "z22", "P11", "P12", "P21", "P22", "R part", "ρcd J/m²K"]
    for j, lab in enumerate(labels):
        c = ws.cell(row=r0 + 1, column=j + 1, value=lab); c.font = F_NOTE
    rows = []
    r = r0 + 2
    # Rsi
    rows.append((r, "Rsi", "res", None, None, f"={rsi}")); r += 1
    for i in range(1, N_SLOTS + 1):
        a, b = mat_col(i)
        rows.append((r, f"Layer {i}", "slot", f"INDEX(Library!${a}${LIB_FIRST}:${a}${LIB_LAST},{idx})", f"INDEX(Library!${b}${LIB_FIRST}:${b}${LIB_LAST},{idx})", None)); r += 1
    rows.append((r, "Far-side layer", "back", None, None, None)); r += 1
    rows.append((r, "Rse", "res", None, None, f'=IF({back}="ext",0.04,IF({back}="room",{rsi},0))')); r += 1
    first, last = rows[0][0], rows[-1][0]
    for rr, lab, kind, matf, valf, rf in rows:
        put(ws, f"A{rr}", lab)
        if kind == "res":
            put(ws, f"B{rr}", "res"); put(ws, f"H{rr}", rf, fmt="0.000")
        elif kind == "slot":
            put(ws, f"C{rr}", f'=IF({matf}="","",{matf})')
            put(ws, f"B{rr}", f'=IF(C{rr}="","none",IF(C{rr}="Cavity","res","mass"))')
            put(ws, f"D{rr}", f'=IF(B{rr}="mass",{valf}/1000,0)', fmt="0.0000")
            put(ws, f"H{rr}", f'=IF(B{rr}="res",{valf},0)', fmt="0.000")
        else:  # far side: soil for ground floors, very large R for adiabatic
            put(ws, f"C{rr}", f'=IF({back}="ground","clay soil","")')
            put(ws, f"B{rr}", f'=IF({back}="ground","mass",IF({back}="adiabatic","res","none"))')
            put(ws, f"D{rr}", f'=IF(B{rr}="mass",1,0)', fmt="0.0000")
            put(ws, f"H{rr}", f'=IF(B{rr}="res",1000000,0)', fmt="0")
        if kind != "res":
            for cc, k in (("E", "B"), ("F", "C"), ("G", "D")):
                put(ws, f"{cc}{rr}", f'=IF(B{rr}="mass",INDEX(Materials!${k}${MAT_FIRST}:${k}${MAT_LAST},MATCH(C{rr},{MATR},0)),0)')
        else:
            put(ws, f"B{rr}", "res")
        put(ws, f"I{rr}", f'=IF(B{rr}="mass",SQRT(E{rr}*{Tsec}/(PI()*F{rr}*G{rr})),0)', fmt="0.0000")
        put(ws, f"J{rr}", f'=IF(B{rr}="mass",D{rr}/I{rr},0)', fmt="0.000")
        x = f"J{rr}"
        sh, chh, co, si = f"SINH({x})", f"COSH({x})", f"COS({x})", f"SIN({x})"
        put(ws, f"K{rr}", f'=IF(B{rr}="mass",COMPLEX({chh}*{co},{sh}*{si}),COMPLEX(1,0))')
        put(ws, f"L{rr}", f'=IF(B{rr}="mass",COMPLEX(-I{rr}/(2*E{rr})*({sh}*{co}+{chh}*{si}),-I{rr}/(2*E{rr})*({chh}*{si}-{sh}*{co})),COMPLEX(-H{rr},0))')
        put(ws, f"M{rr}", f'=IF(B{rr}="mass",COMPLEX(-E{rr}/I{rr}*({sh}*{co}-{chh}*{si}),-E{rr}/I{rr}*({sh}*{co}+{chh}*{si})),COMPLEX(0,0))')
        put(ws, f"N{rr}", f"=K{rr}")
        if rr == first:
            for a_, b_ in zip("OPQR", "KLMN"):
                put(ws, f"{a_}{rr}", f"={b_}{rr}")
        else:
            p = rr - 1
            put(ws, f"O{rr}", f"=IMSUM(IMPRODUCT(K{rr},O{p}),IMPRODUCT(L{rr},Q{p}))")
            put(ws, f"P{rr}", f"=IMSUM(IMPRODUCT(K{rr},P{p}),IMPRODUCT(L{rr},R{p}))")
            put(ws, f"Q{rr}", f"=IMSUM(IMPRODUCT(M{rr},O{p}),IMPRODUCT(N{rr},Q{p}))")
            put(ws, f"R{rr}", f"=IMSUM(IMPRODUCT(M{rr},P{p}),IMPRODUCT(N{rr},R{p}))")
        put(ws, f"S{rr}", f'=IF(B{rr}="mass",D{rr}/E{rr},H{rr})', fmt="0.000")
        put(ws, f"T{rr}", f'=IF(B{rr}="mass",F{rr}*G{rr}*D{rr},0)', fmt="0")
    # outputs
    o = last + 1
    Z11, Z12 = f"O{last}", f"P{last}"
    labels = [("U W/m²K", f'=IF(OR({back}="ext",{back}="ground"),1/SUM(S{first}:S{last}),"–")', "0.00"),
              ("|Y| W/m²K", f"=IMABS(IMDIV(IMPRODUCT({Z11},COMPLEX(-1,0)),{Z12}))", "0.00"),
              ("Y lead h", f"=IMARGUMENT(IMDIV(IMPRODUCT({Z11},COMPLEX(-1,0)),{Z12}))*{period_ref}/(2*PI())", "0.0"),
              ("f", f'=IF({back}="ext",IMABS(IMDIV(COMPLEX(-1,0),{Z12}))*SUM(S{first}:S{last}),"–")', "0.00"),
              ("lag h", f'=IF({back}="ext",MOD(-IMARGUMENT(IMDIV(COMPLEX(-1,0),{Z12}))*{period_ref}/(2*PI()),24),"–")', "0.0"),
              ("κ₁ kJ/m²K", f"={Tsec}/(2*PI())*IMABS(IMDIV(IMSUB({Z11},COMPLEX(1,0)),{Z12}))/1000", "0"),
              ("κ total kJ/m²K", f"=SUM(T{first + 1}:T{first + N_SLOTS})/1000", "0"),
              ("Rc m²K/W", f"=SUM(S{first + 1}:S{first + N_SLOTS})", "0.000")]
    refs = {}
    for j, (lab, f, fmt) in enumerate(labels):
        c0 = 1 + 2 * j
        put(ws, f"{col(c0)}{o}", lab, F_NOTE)
        out(ws, f"{col(c0 + 1)}{o}", f, fmt)
        refs[lab] = f"{col(c0 + 1)}{o}"
    refs.update(first=first, last=last, back=back, rsi=rsi, idx=idx, end=o)
    return refs


def fivenode_block(ws, r0, iso, cls_ref, period_ref, title):
    """Five-node chain: Rsi, C5, Rc/6, C4, Rc/3, C3, Rc/3, C2, Rc/6, C1, far-side layer, Rse."""
    put(ws, f"A{r0}", title, F_B)
    w = f"(2*PI()/({period_ref}*3600))"
    Rc, kap = iso["Rc m²K/W"], f"({iso['κ total kJ/m²K']}*1000)"
    frac = lambda n: f"INDEX(Classes!$B$5:$F$9,MATCH({cls_ref},Classes!$A$5:$A$9,0),{n})"  # n = 1 ext .. 5 int
    seq = [("Rsi", "res", f"={iso['rsi']}"), ("node 5 (room side)", "cap", f"={frac(5)}*{kap}"),
           ("6/Rc", "res", f"={Rc}/6"), ("node 4", "cap", f"={frac(4)}*{kap}"), ("3/Rc", "res", f"={Rc}/3"),
           ("node 3", "cap", f"={frac(3)}*{kap}"), ("3/Rc", "res", f"={Rc}/3"), ("node 2", "cap", f"={frac(2)}*{kap}"),
           ("6/Rc", "res", f"={Rc}/6"), ("node 1 (outer)", "cap", f"={frac(1)}*{kap}"), ("far-side layer", "copy", iso["last"] - 1),
           ("Rse", "copy", iso["last"])]
    for j, lab in enumerate(["", "Kind", "R or C"] + [""] * 7 + ["z11", "z12", "z21", "z22", "P11", "P12", "P21", "P22"]):
        c = ws.cell(row=r0 + 1, column=j + 1, value=lab); c.font = F_NOTE
    r = r0 + 2; first = r
    for lab, kind, v in seq:
        put(ws, f"A{r}", lab)
        if kind == "copy":
            put(ws, f"B{r}", "as ISO")
            for cc in "KLMN":
                put(ws, f"{cc}{r}", f"={cc}{v}")
        else:
            put(ws, f"B{r}", kind); put(ws, f"C{r}", v, fmt="0.000")
            if kind == "res":
                put(ws, f"K{r}", "=COMPLEX(1,0)"); put(ws, f"L{r}", f"=COMPLEX(-C{r},0)"); put(ws, f"M{r}", "=COMPLEX(0,0)"); put(ws, f"N{r}", "=COMPLEX(1,0)")
            else:
                put(ws, f"K{r}", "=COMPLEX(1,0)"); put(ws, f"L{r}", "=COMPLEX(0,0)"); put(ws, f"M{r}", f"=COMPLEX(0,-{w}*C{r})"); put(ws, f"N{r}", "=COMPLEX(1,0)")
        if r == first:
            for a_, b_ in zip("OPQR", "KLMN"):
                put(ws, f"{a_}{r}", f"={b_}{r}")
        else:
            p = r - 1
            put(ws, f"O{r}", f"=IMSUM(IMPRODUCT(K{r},O{p}),IMPRODUCT(L{r},Q{p}))")
            put(ws, f"P{r}", f"=IMSUM(IMPRODUCT(K{r},P{p}),IMPRODUCT(L{r},R{p}))")
            put(ws, f"Q{r}", f"=IMSUM(IMPRODUCT(M{r},O{p}),IMPRODUCT(N{r},Q{p}))")
            put(ws, f"R{r}", f"=IMSUM(IMPRODUCT(M{r},P{p}),IMPRODUCT(N{r},R{p}))")
        r += 1
    last = r - 1
    put(ws, f"A{r}", "|Y| five-node W/m²K", F_NOTE)
    out(ws, f"B{r}", f"=IMABS(IMDIV(IMPRODUCT(O{last},COMPLEX(-1,0)),P{last}))", "0.00")
    return f"B{r}", r


def classes_sheet(wb):
    ws = sheet(wb, "Classes", {"A": 10, "B": 12, "C": 12, "D": 12, "E": 12, "F": 12})
    put(ws, "A1", "Five-node mass classes (ISO 52016-1 / HEM)", F_T)
    note(ws, "A2", "Share of the element's total areal heat capacity placed on each node. Node 1 is the outer node, node 5 the room-side surface node.")
    header(ws, 4, 1, ["Class", "Node 1 (outer)", "Node 2", "Node 3", "Node 4", "Node 5 (room)"])
    for i, (k, v) in enumerate(CLASSES.items()):
        put(ws, f"A{5 + i}", k, bold=True)
        for j, f in enumerate(v):
            inp(ws, f"{col(2 + j)}{5 + i}", f, "0.000")
    note(ws, "A11", "I internal, E external, IE internal and external, D distributed, M middle. Conductances between nodes are 6/Rc, 3/Rc, 3/Rc, 6/Rc.")


def calculator_sheet(wb):
    ws = sheet(wb, "Calculator", {c: 13 for c in "ABCDEFGHIJKLMNOPQRSTUVWXY"})
    ws.column_dimensions["A"].width = 20
    put(ws, "A1", "Build-up calculator", F_T)
    note(ws, "A2", "Choose a build-up (or 'Custom' to use your own row on the Library sheet) and a period. Everything below recalculates.")
    put(ws, "A4", "Build-up key", bold=True); inp(ws, "B4", "W2")
    dv_list(ws, "B4", f"=Library!$A${LIB_FIRST}:$A${LIB_LAST}")
    put(ws, "A5", "Period, hours", bold=True); inp(ws, "B5", 24)
    put(ws, "A6", "Five-node class", bold=True); inp(ws, "B6", "rule")
    dv_list(ws, "B6", '"rule,I,E,IE,D,M"')
    put(ws, "C4", f"=INDEX(Library!$B${LIB_FIRST}:$B${LIB_LAST},MATCH(B4,Library!$A${LIB_FIRST}:$A${LIB_LAST},0))", F_B)
    put(ws, "C6", f'=IF(B6="rule",INDEX(Library!$F${LIB_FIRST}:$F${LIB_LAST},MATCH(B4,Library!$A${LIB_FIRST}:$A${LIB_LAST},0)),B6)')
    note(ws, "D6", "← class used below")
    iso = chain_block(ws, 9, "$B$4", "$B$5", "ISO 13786 matrix chain (room side first)")
    r = iso["end"] + 3
    y5, r = fivenode_block(ws, r, iso, "$C$6", "$B$5", "Five-node chain for the chosen class")
    r += 2
    put(ws, f"A{r}", "Five-node |Y| by class, same period", F_B)
    header(ws, r + 1, 1, ["Class", "|Y| W/m²K", "÷ ISO"])
    # all five classes on a hidden-ish block to the right
    rr = r + 2
    blocks = {}
    base = r + 10
    for k, c in enumerate(CLASSES):
        put(ws, f"A{rr + k}", c, bold=True)
        ref, end = fivenode_block(ws, base, iso, f'"{c}"', "$B$5", f"Five-node chain, class {c}")
        out(ws, f"B{rr + k}", f"={ref}", "0.00"); out(ws, f"C{rr + k}", f"=B{rr + k}/{iso['|Y| W/m²K']}", "0.00")
        base = end + 3
    out(ws, f"B{rr + 5}", f"={iso['|Y| W/m²K']}", "0.00"); put(ws, f"A{rr + 5}", "ISO 13786", bold=True)
    ws.conditional_formatting.add(f"C{rr}:C{rr + 4}", ColorScaleRule(start_type="num", start_value=0.5, start_color="5B9BD5", mid_type="num", mid_value=1, mid_color="FFFFFF", end_type="num", end_value=2, end_color="E46C0A"))
    # summary box at the top
    put(ws, "F4", "Results", bold=True)
    for j, (lab, ref) in enumerate([(k, iso[k]) for k in ("U W/m²K", "|Y| W/m²K", "Y lead h", "f", "lag h", "κ₁ kJ/m²K", "κ total kJ/m²K")] + [("Five-node |Y| (class)", y5)]):
        put(ws, f"{col(6 + j)}5", lab, F_NOTE, align="wrap")
        out(ws, f"{col(6 + j)}6", f"={ref}", ws[ref].number_format, bold=True)
    ws.row_dimensions[5].height = 28
    ws.freeze_panes = "A8"
    return ws


def all_sheet(wb):
    """Every library build-up computed live, one block per build-up."""
    eng = sheet(wb, "Engine", {c: 12 for c in "ABCDEFGHIJKLMNOPQRSTUVWXY"})
    put(eng, "A1", "Calculation blocks for the Library results sheet (one ISO chain and one five-node chain per build-up)", F_B)
    put(eng, "A2", "Period, hours", bold=True); inp(eng, "B2", 24)
    res = sheet(wb, "Library results", {"A": 8, "B": 52, "C": 8, "D": 9, "E": 9, "F": 9, "G": 9, "H": 9, "I": 9, "J": 9, "K": 11, "L": 11, "M": 11, "N": 11})
    put(res, "A1", "All build-ups at the period on the Engine sheet (24 h)", F_T)
    note(res, "A2", "Live formulas. The last two columns are the Python model's values for the same build-up, for checking.")
    header(res, 4, 1, ["Key", "Build-up", "Class", "U", "|Y|", "Y lead h", "f", "lag h", "κ₁", "κ total", "Five-node |Y|", "5-node ÷ ISO", "Python |Y|", "Python 5-node |Y|"])
    py = {r["key"]: r for r in csv.DictReader(open(os.path.join(MODELS, "results", "buildups.csv")))}
    r0 = 5
    keys = list(LIB) + ["Custom"]
    for i, key in enumerate(keys):
        put(eng, f"A{r0}", key, bold=True)
        iso = chain_block(eng, r0 + 1, f"$A${r0}", "$B$2", f"{key}: ISO 13786")
        y5, end = fivenode_block(eng, iso["end"] + 2, iso, f"Engine!$C${r0}", "$B$2", f"{key}: five-node")
        put(eng, f"B{r0}", "class", F_NOTE); put(eng, f"C{r0}", f"=INDEX(Library!$F${LIB_FIRST}:$F${LIB_LAST},MATCH(A{r0},Library!$A${LIB_FIRST}:$A${LIB_LAST},0))")
        rr = 5 + i
        put(res, f"A{rr}", key, bold=True)
        put(res, f"B{rr}", f"=INDEX(Library!$B${LIB_FIRST}:$B${LIB_LAST},MATCH(A{rr},Library!$A${LIB_FIRST}:$A${LIB_LAST},0))")
        put(res, f"C{rr}", f"=Engine!C{r0}")
        for j, k in enumerate(("U W/m²K", "|Y| W/m²K", "Y lead h", "f", "lag h", "κ₁ kJ/m²K", "κ total kJ/m²K")):
            out(res, f"{col(4 + j)}{rr}", f"=Engine!{iso[k]}", eng[iso[k]].number_format)
        out(res, f"K{rr}", f"=Engine!{y5}", "0.00")
        out(res, f"L{rr}", f"=K{rr}/E{rr}", "0.00")
        if key in py:
            put(res, f"M{rr}", round(float(py[key]["Y"]), 3), fmt="0.00")
            put(res, f"N{rr}", round(float(py[key]["Y_5n"]), 3), fmt="0.00")
        r0 = end + 3
    res.conditional_formatting.add(f"L5:L{4 + len(keys)}", ColorScaleRule(start_type="num", start_value=0.3, start_color="5B9BD5", mid_type="num", mid_value=1, mid_color="FFFFFF", end_type="num", end_value=2.5, end_color="E46C0A"))
    res.freeze_panes = "C5"
    return res


def build(path):
    wb = new_wb()
    readme(wb, "Build-up Explorer (thermal mass toolkit, model 1)", [
        "What it does:",
        "Computes the ISO 13786 dynamic thermal properties of a layered wall, floor or roof (U, admittance |Y| and its time lead, decrement factor f and time lag, internal areal heat capacity κ₁) and compares them with the ISO 52016-1 / HEM five-node lump for each mass class.",
        "Everything is a live formula. The heat-transfer matrices are complex numbers, built with COMPLEX, IMPRODUCT, IMSUM and IMDIV and multiplied layer by layer from the room side out.",
        "",
        "How to use it:",
        "1. Calculator: pick a build-up key in B4 (or 'Custom'), a period in B5 (24 h for the daily cycle), and optionally a five-node class in B6.",
        "2. Library: the 29 build-ups from the Python model, plus a Custom row at the bottom you can edit. Layer values are mm, or m²K/W for a 'Cavity'.",
        "3. Library results: every build-up at once, with the Python model's values alongside for checking.",
        "4. Materials and Classes: edit material properties or the class node shares; everything follows.",
        "",
        "Far side: 'ext' adds Rse = 0.04; 'ground' adds 1 m of clay soil; 'adiabatic' (half a partition) adds a very large resistance; 'room' (an intermediate floor) adds Rsi on the far face.",
        "Companion files: the web app Build-up Explorer and models/buildups.py in gwilma/thermal_mass_in_buildings give the same numbers.",
    ])
    calculator_sheet(wb)
    all_sheet(wb)
    library_sheet(wb)
    materials_sheet(wb)
    classes_sheet(wb)
    wb.move_sheet("Calculator", offset=-(len(wb.sheetnames) - 2))
    wb.save(path)


if __name__ == "__main__":
    os.makedirs(os.path.join(HERE, "out"), exist_ok=True)
    build(os.path.join(HERE, "out", "buildup-explorer.xlsx"))
    print("written")
