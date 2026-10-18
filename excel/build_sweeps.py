"""Workbook 4: Mass Sweeps. Results of the Python parameter sweeps, with live
derived columns, an interpolator on the mass x night-vent grid, and charts.

The sweeps themselves (hundreds of full-year runs) stay in Python and the web
app; this workbook holds their output so it can be charted and compared.
"""
import csv, json, os
from openpyxl.chart import BarChart, LineChart, Reference, ScatterChart, Series
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment
from xlcommon import (F, F_B, F_NOTE, FILL_OUT, HERE, MODELS, col, dv_list, header,
                      inp, new_wb, note, out, put, readme, sheet)

RES = os.path.join(MODELS, "results")
S = json.load(open(os.path.join(RES, "sweeps.json")))
PK = list(csv.DictReader(open(os.path.join(RES, "packages.csv"))))

wb = new_wb()
readme(wb, "Mass Sweeps: results workbook", [
    "What this is:",
    "The results of model 4, the parameter sweeps, run in Python (models/sweeps.py) on the full fine-mesh room model with Heathrow TMYx 2011-2025 weather. The same sweeps run live in the Mass Sweeps web app.",
    "Each sweep is hundreds of full-year hourly runs, which is too much for a spreadsheet, so the numbers here are pasted results. Everything derived from them (percentage savings, marginal gains, the interpolator) is a live formula.",
    "",
    "Base case:",
    "The 'Timber frame, screed ground floor' house (90 m², two storeys, 18 m² glazing, front faces south). An exposed dense layer of thickness t is added to the ground floor, the intermediate floor (a slab of thickness t), the partitions (2t thick) and the party wall. t = 0 swaps in board and chipboard.",
    "Hot year: TMYx +1.5 K June to August plus a +4 K heatwave bump over 14 to 21 July. 2050s: hot year +2.5 K summer, +1.5 K otherwise.",
    "",
    "Sheets:",
    "Diminishing returns: sweep A, nine thicknesses from 0 to 300 mm, with savings against t = 0 and the marginal gain per extra millimetre.",
    "Mass x night vent: sweep B, a 6 x 7 grid of night ventilation rate against exposed thickness, as heatmaps, plus an interpolator you can type into.",
    "Morris: sweep C, Morris screening of eight inputs for four outputs. μ* ranks how much each input moves the output over its full range; σ shows how much that effect depends on the other inputs.",
    "Packages: the 14 construction packages from the House Simulator, with the full model and the simpler 5R1C model side by side.",
    "",
    "Metrics:",
    "DH26: degree-hours above 26 °C operative, May to September. TM59 (b) style: night hours (23-07) above 26 °C, limit 32. TM59 (a) style: % of occupied hours 1 K or more over the adaptive limit, limit 3%. These are single-zone whole-house figures, so they indicate rather than test TM59 compliance.",
    "κ₁: areal heat capacity on the room side of each element (ISO 13786, 24 h), summed over all elements and divided by floor area.",
])

# ---------------------------------------------------------------- sweep A
ws = sheet(wb, "Diminishing returns", {"A": 9, **{col(c): 11.5 for c in range(2, 24)}})
put(ws, "A1", "Sweep A: exposed thickness t", F_B)
note(ws, "A2", "Pasted results from models/sweeps.py (blue). Derived columns on the right are formulas.")
cols = [("t mm", None, "0"), ("κ₁ kJ/m²K per m² floor", "kappa1", "0"), ("κ_tot kJ/m²K per m² floor", "kappa_tot", "0"),
        ("Heat kWh/m² (typical)", "heat", "0.0"), ("Peak heat kW", "peak", "0.00"),
        ("DH26 hot, day purge", "hot_day_dh26", "0"), ("DH26 hot, + night vent", "hot_nv_dh26", "0"),
        ("DH26 2050s, + night vent", "hot2050_nv_dh26", "0"), ("(b) night h hot + nv", "hot_nv_b", "0"),
        ("(a) % hot + nv", "hot_nv_a", "0.0"), ("T_op max hot + nv °C", "hot_nv_max", "0.0"),
        ("July swing hot + nv K", "hot_nv_swing", "0.0")]
derived = ["Heat saving vs t=0", "Peak cut vs t=0", "DH26 cut (day) vs t=0", "DH26 cut (+nv) vs t=0",
           "Marginal DH26 (+nv) per mm"]
header(ws, 4, 1, [c[0] for c in cols] + [""] + derived)
r0 = 5
for i, rec in enumerate(S["A"]):
    r = r0 + i
    for j, (lab, key, fmt) in enumerate(cols):
        v = round(rec["t"] * 1000) if key is None else rec[key]
        inp(ws, f"{col(j + 1)}{r}", v, fmt)
    out(ws, f"N{r}", f"=1-D{r}/$D${r0}", "0%")
    out(ws, f"O{r}", f"=1-E{r}/$E${r0}", "0%")
    out(ws, f"P{r}", f"=1-F{r}/$F${r0}", "0%")
    out(ws, f"Q{r}", f"=1-G{r}/$G${r0}", "0%")
    out(ws, f"R{r}", f'=IF(A{r}=$A${r0},"",(G{r - 1}-G{r})/(A{r}-A{r - 1}))' if i else '=""', "0.0")
rN = r0 + len(S["A"]) - 1
note(ws, f"A{rN + 2}", "Marginal DH26 per mm: degree-hours saved by each extra millimetre of exposed dense layer, between this row and the one above. It falls as t passes the daily penetration depth (about 100 to 150 mm in dense concrete).")
note(ws, f"A{rN + 3}", "κ₁ peaks near t = 100 to 150 mm and then falls slightly: beyond the penetration depth extra mass adds total heat capacity but not daily admittance.")

def line(title, ytitle, ycols, anchor, rmin=r0, rmax=rN, ws=ws, w=17, h=8.5):
    ch = ScatterChart(); ch.title = title; ch.style = 13
    ch.x_axis.title = "Exposed thickness t, mm"; ch.y_axis.title = ytitle
    ch.height, ch.width = h, w
    xs = Reference(ws, min_col=1, min_row=rmin, max_row=rmax)
    for c in ycols:
        s = Series(Reference(ws, min_col=c, min_row=rmin - 1, max_row=rmax), xs, title_from_data=True)
        s.marker.symbol = "circle"; s.marker.size = 5
        ch.series.append(s)
    ch.x_axis.delete = False; ch.y_axis.delete = False
    ch.legend.position = "b"
    ws.add_chart(ch, anchor)

line("Overheating falls steadily with exposed mass", "Degree-hours > 26 °C", [6, 7, 8], f"A{rN + 5}")
line("Heating savings flatten by about 50 mm", "kWh/m² or kW", [4, 5], f"J{rN + 5}")
line("Daily admittance saturates; total capacity keeps rising", "kJ/m²K per m² floor", [2, 3], f"A{rN + 23}")
line("Peak temperature and July swing", "°C or K", [11, 12], f"J{rN + 23}")
ws.freeze_panes = "B5"

# ---------------------------------------------------------------- sweep B
ws = sheet(wb, "Mass x night vent", {"A": 22, **{col(c): 10 for c in range(2, 12)}})
B = S["B"]; ts = [round(t * 1000) for t in B["t"]]; ach = B["ach"]
nt, na = len(ts), len(ach)
put(ws, "A1", "Sweep B: night ventilation against exposed thickness (hot year)", F_B)

def grid(r, title, data, fmt, note_text):
    put(ws, f"A{r}", title, F_B)
    header(ws, r + 1, 1, ["Night ACH \\ t mm"] + ts)
    for i, a in enumerate(ach):
        inp(ws, f"A{r + 2 + i}", a, "0.0")
        for j in range(nt):
            inp(ws, f"{col(j + 2)}{r + 2 + i}", data[i][j], fmt)
    for j in range(nt):  # t values as numbers for the interpolator
        ws[f"{col(j + 2)}{r + 1}"].number_format = "0"
    rng = f"B{r + 2}:{col(nt + 1)}{r + 1 + na}"
    ws.conditional_formatting.add(rng, ColorScaleRule(start_type="min", start_color="5B8DB8",
                                                      mid_type="percentile", mid_value=50, mid_color="F6F1E3",
                                                      end_type="max", end_color="C8553D"))
    note(ws, f"A{r + 2 + na}", note_text)
    return r + 2, rng

gA, rngA = grid(3, "Degree-hours above 26 °C, May to September", B["dh26"], "0",
                "Reading down a column: ventilation at fixed mass. Reading across a row: mass at fixed ventilation. Mass does most when it is flushed at night.")
gB, rngB = grid(gA + na + 2, "Night hours above 26 °C (TM59 (b) style, limit 32)", B["night"], "0",
                "Without night ventilation, mass holds heat into the night: (b) gets worse from t = 0 to 50 mm before improving.")

r = gB + na + 2
put(ws, f"A{r}", "Relative benefit of mass: DH26 cut from t = 0 to t, at each ventilation rate", F_B)
header(ws, r + 1, 1, ["Night ACH \\ t mm"] + ts)
for i in range(na):
    put(ws, f"A{r + 2 + i}", f"=A{gA + i}", fmt="0.0")
    for j in range(nt):
        out(ws, f"{col(j + 2)}{r + 2 + i}", f"=1-{col(j + 2)}{gA + i}/$B{gA + i}", "0%")
ws.conditional_formatting.add(f"B{r + 2}:{col(nt + 1)}{r + 1 + na}",
                              ColorScaleRule(start_type="min", start_color="F6F1E3", end_type="max", end_color="3E7C59"))
rel = r

# interpolator
r = rel + na + 3
put(ws, f"A{r}", "Interpolator (bilinear within the grid)", F_B)
put(ws, f"A{r + 1}", "Exposed thickness t, mm"); inp(ws, f"B{r + 1}", 80, "0")
put(ws, f"A{r + 2}", "Night ventilation, ACH"); inp(ws, f"B{r + 2}", 1.5, "0.0")
tr = f"$B${gA - 1}:${col(nt + 1)}${gA - 1}"
ar = f"$A${gA}:$A${gA + na - 1}"
T, A_ = f"MAX({ts[0]},MIN({ts[-1]},B{r + 1}))", f"MAX({ach[0]},MIN({ach[-1]},B{r + 2}))"
put(ws, f"D{r + 1}", "col i"); out(ws, f"E{r + 1}", f"=MIN(MATCH({T},{tr},1),{nt - 1})", "0")
put(ws, f"D{r + 2}", "row k"); out(ws, f"E{r + 2}", f"=MIN(MATCH({A_},{ar},1),{na - 1})", "0")
put(ws, f"D{r + 3}", "u"); out(ws, f"E{r + 3}", f"=({T}-INDEX({tr},E{r + 1}))/(INDEX({tr},E{r + 1}+1)-INDEX({tr},E{r + 1}))", "0.00")
put(ws, f"D{r + 4}", "v"); out(ws, f"E{r + 4}", f"=({A_}-INDEX({ar},E{r + 2}))/(INDEX({ar},E{r + 2}+1)-INDEX({ar},E{r + 2}))", "0.00")

def bil(g):
    G = f"$B${g}:${col(nt + 1)}${g + na - 1}"
    i, k, u, v = f"$E${r + 1}", f"$E${r + 2}", f"$E${r + 3}", f"$E${r + 4}"
    return (f"=(1-{u})*(1-{v})*INDEX({G},{k},{i})+{u}*(1-{v})*INDEX({G},{k},{i}+1)"
            f"+(1-{u})*{v}*INDEX({G},{k}+1,{i})+{u}*{v}*INDEX({G},{k}+1,{i}+1)")
put(ws, f"A{r + 3}", "DH26 hot year", F_B); out(ws, f"B{r + 3}", bil(gA), "0", bold=True)
put(ws, f"A{r + 4}", "Night hours > 26 °C", F_B); out(ws, f"B{r + 4}", bil(gB), "0", bold=True)
note(ws, f"A{r + 6}", f"Inputs outside the grid ({ts[0]}-{ts[-1]} mm, {ach[0]}-{ach[-1]} ACH) are clamped to its edge. Interpolation is linear, so between coarse grid points it is approximate.")

# chart: DH26 vs t for each ACH
ch = ScatterChart(); ch.title = "DH26 against exposed thickness, one line per night ventilation rate"
ch.style = 13; ch.height, ch.width = 9, 18
ch.x_axis.title = "Exposed thickness t, mm"; ch.y_axis.title = "Degree-hours > 26 °C"
xs = Reference(ws, min_col=2, max_col=nt + 1, min_row=gA - 1)
for i in range(na):
    s = Series(Reference(ws, min_col=2, max_col=nt + 1, min_row=gA + i), xs, title=f"{ach[i]} ACH")
    s.marker.symbol = "circle"; s.marker.size = 5
    ch.series.append(s)
ch.x_axis.delete = False; ch.y_axis.delete = False; ch.legend.position = "r"
ws.add_chart(ch, "J3")

# ---------------------------------------------------------------- Morris
ws = sheet(wb, "Morris", {"A": 26, "B": 11, "C": 11, "D": 11, "E": 3})
put(ws, "A1", "Sweep C: Morris screening (10 trajectories, 4 levels, 8 inputs, 90 runs per weather)", F_B)
note(ws, "A2", "μ*: mean absolute elementary effect over the full input range (bigger = more influential). μ: signed mean (sign = direction). σ: spread (large = the effect depends on other inputs).")
note(ws, "A3", "Ranges: exposed mass t 0-0.3 m; glazing × 0.5-1.5; blind factor 0.2-1.0; day purge 0-6 ACH; night vent 0-4 ACH; background 0.3-1.0 ACH; gains × 0.6-1.4; wall insulation 0.09-0.24 m mineral wool.")
titles = {"dh26": "Degree-hours > 26 °C, hot year", "tm59b_h": "Night hours > 26 °C, hot year",
          "heat_kwh_m2": "Space heating kWh/m², typical year", "heat_peak_kw": "Peak heating kW, typical year"}
r = 5
for k, rows in S["morris"].items():
    rows = sorted(rows, key=lambda x: -x["mu_star"])
    put(ws, f"A{r}", titles.get(k, k), F_B)
    header(ws, r + 1, 1, ["Input", "μ*", "μ (signed)", "σ"])
    fmt = "0.00" if k == "heat_peak_kw" else ("0.0" if k == "heat_kwh_m2" else "0")
    for i, x in enumerate(rows):
        put(ws, f"A{r + 2 + i}", x["name"])
        inp(ws, f"B{r + 2 + i}", x["mu_star"], fmt); inp(ws, f"C{r + 2 + i}", x["mu"], fmt); inp(ws, f"D{r + 2 + i}", x["sigma"], fmt)
    n = len(rows)
    put(ws, f"A{r + 2 + n}", "Mass rank (1 = most influential)")
    out(ws, f"B{r + 2 + n}", f'=RANK(INDEX(B{r + 2}:B{r + 1 + n},MATCH("exposed mass t (m)",A{r + 2}:A{r + 1 + n},0)),B{r + 2}:B{r + 1 + n})', "0")
    bc = BarChart(); bc.type = "bar"; bc.style = 10; bc.title = titles.get(k, k) + ": μ*"
    bc.add_data(Reference(ws, min_col=2, min_row=r + 1, max_row=r + 1 + n), titles_from_data=True)
    bc.set_categories(Reference(ws, min_col=1, min_row=r + 2, max_row=r + 1 + n))
    bc.x_axis.scaling.orientation = "maxMin"; bc.legend = None
    bc.height, bc.width = 6.5, 13
    bc.x_axis.delete = False; bc.y_axis.delete = False
    ws.add_chart(bc, f"F{r}")
    r += n + 5
note(ws, f"A{r}", "Read-across: shading (blinds, glazing area) and day purge dominate summer outcomes; mass ranks mid-table for overheating but second (after background ventilation) for peak heating load, where it smooths the morning reheat.")

# ---------------------------------------------------------------- Packages
ws = sheet(wb, "Packages", {"A": 46, **{col(c): 11 for c in range(2, 20)}})
put(ws, "A1", "Construction packages: full fine-mesh model against 5R1C", F_B)
note(ws, "A2", "From models/packages.py with default House Simulator inputs. 'Full' keeps every layer of every element (fine mesh); 5R1C lumps the whole house into one mass node (ISO 13790).")
pc = [("κ₁ per m² floor", "kappa1", "0"), ("H fabric W/K", "H_fabric", "0"),
      ("Heat kWh/m² full", "heat_kwh_m2", "0.0"), ("Heat kWh/m² 5R1C", "r5_heat", "0.0"),
      ("Peak kW full", "heat_peak_kw", "0.00"), ("Peak kW 5R1C", "r5_peak", "0.00"),
      ("DH26 hot full", "hot_day_dh26", "0"), ("DH26 hot 5R1C", "r5_day_dh26", "0"),
      ("DH26 hot+nv full", "hot_nv_dh26", "0"), ("DH26 hot+nv 5R1C", "r5_nv_dh26", "0"),
      ("(b) h hot+nv full", "hot_nv_tm59b_h", "0"), ("(b) h hot+nv 5R1C", "r5_nv_b", "0"),
      ("DH26 2050s+nv full", "hot2050_nv_dh26", "0"), ("July swing hot+nv K", "hot_nv_jul_swing", "0.0")]
header(ws, 4, 1, ["Package"] + [c[0] for c in pc] + ["5R1C/full heat", "5R1C/full DH26+nv"])
rows = [p for p in PK if p["mode"] == "fd"]
for i, p in enumerate(rows):
    rr = 5 + i
    put(ws, f"A{rr}", p["package"])
    for j, (lab, key, fmt) in enumerate(pc):
        inp(ws, f"{col(j + 2)}{rr}", float(p[key]), fmt)
    out(ws, f"P{rr}", f"=E{rr}/D{rr}", "0.00")
    out(ws, f"Q{rr}", f"=K{rr}/J{rr}", "0.00")
pN = 4 + len(rows)
for c in ("D", "H", "J"):
    ws.conditional_formatting.add(f"{c}5:{c}{pN}", ColorScaleRule(start_type="min", start_color="5B8DB8",
                                  mid_type="percentile", mid_value=50, mid_color="F6F1E3", end_type="max", end_color="C8553D"))
note(ws, f"A{pN + 2}", "The two models agree on heating to within about 10% (5R1C a little higher). In summer 5R1C under-rewards deep, well-ventilated mass (heavyweight with night vent: 1208 against 683 degree-hours) because one node cannot hold a night's coolth separately from the day's surface swing.")
note(ws, f"A{pN + 3}", "The 'Solid brick, uninsulated' row is off-scale for heating and is left in for reference.")
bc = BarChart(); bc.type = "bar"; bc.style = 10; bc.title = "DH26 hot year with night ventilation"
bc.add_data(Reference(ws, min_col=10, max_col=11, min_row=4, max_row=pN), titles_from_data=True)
bc.set_categories(Reference(ws, min_col=1, min_row=5, max_row=pN))
bc.x_axis.scaling.orientation = "maxMin"; bc.height, bc.width = 11, 20; bc.legend.position = "b"
bc.x_axis.delete = False; bc.y_axis.delete = False
ws.add_chart(bc, f"A{pN + 5}")
bc = BarChart(); bc.type = "bar"; bc.style = 10; bc.title = "Space heating, typical year, kWh/m²"
bc.add_data(Reference(ws, min_col=4, max_col=5, min_row=4, max_row=pN), titles_from_data=True)
bc.set_categories(Reference(ws, min_col=1, min_row=5, max_row=pN))
bc.x_axis.scaling.orientation = "maxMin"; bc.height, bc.width = 11, 20; bc.legend.position = "b"
bc.x_axis.delete = False; bc.y_axis.delete = False
ws.add_chart(bc, f"I{pN + 5}")
ws.freeze_panes = "B5"

os.makedirs(os.path.join(HERE, "out"), exist_ok=True)
fn = os.path.join(HERE, "out", "mass-sweeps.xlsx")
wb.save(fn)
print("written", fn)
