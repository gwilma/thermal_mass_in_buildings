"""The Excel workbooks on the same reference cases.

Workbooks are edited with openpyxl and recalculated by LibreOffice headless
(the same engine used to check them when they were built), then read back.
The delivered files in excel/out already hold calculated values, which the
parity tests read directly.
"""
import math, os, shutil, subprocess, tempfile
import numpy as np
import pytest
from openpyxl import load_workbook

import room, room5r1c, sweeps
from buildups import LIB, Element, L
from conftest import ROOT

OUT = os.path.join(ROOT, "excel", "out")
pytestmark = pytest.mark.excel


def recalc(src, edit, soffice):
    """Copy a workbook, apply `edit(wb)`, recalculate in LibreOffice, return values."""
    tmp = tempfile.mkdtemp()
    try:
        f = os.path.join(tmp, "in.xlsx")
        wb = load_workbook(os.path.join(OUT, src)); edit(wb); wb.save(f)
        subprocess.run([soffice, "--headless", "--calc", "--convert-to", "xlsx", "--outdir", os.path.join(tmp, "out"), f],
                       check=True, capture_output=True, timeout=600)
        return load_workbook(os.path.join(tmp, "out", "in.xlsx"), data_only=True)
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def _custom(wb, layers, back):
    """Write layers [(material, mm)] into the Library sheet's Custom row."""
    ws = wb["Library"]
    r = next(c.row for c in ws["A"] if c.value == "Custom")
    ws.cell(r, 4).value = back
    for k in range(6):
        ws.cell(r, 7 + 2 * k).value = layers[k][0] if k < len(layers) else None
        ws.cell(r, 8 + 2 * k).value = layers[k][1] if k < len(layers) else None


# --- Build-up Explorer ------------------------------------------------------------
def test_buildup_workbook_annex_d1(soffice):
    """ISO 13786 Annex D case D.1 typed into the Custom row: U 3.56, |Y| 5.70,
    |Y12| = f U 1.83 W/m2K, lag 5.7 h."""
    def edit(wb):
        m = wb["Materials"]
        r = next(c.row for c in m["A"] if c.value == "dense concrete")
        m.cell(r, 2).value, m.cell(r, 3).value, m.cell(r, 4).value = 1.8, 2400, 1000
        _custom(wb, [("dense concrete", 200)], "ext")
        wb["Calculator"]["B4"] = "Custom"
    ws = recalc("buildup-explorer.xlsx", edit, soffice)["Calculator"]
    U, Y, f, lag = ws["F6"].value, ws["G6"].value, ws["I6"].value, ws["J6"].value
    assert U == pytest.approx(3.56, abs=0.005)
    assert Y == pytest.approx(5.70, abs=0.005)
    assert f * U == pytest.approx(1.83, abs=0.005)
    assert lag == pytest.approx(5.7, abs=0.05)


def test_buildup_workbook_library_matches_python():
    ws = load_workbook(os.path.join(OUT, "buildup-explorer.xlsx"), data_only=True)["Library results"]
    seen = 0
    for row in ws.iter_rows(min_row=5, values_only=True):
        key = row[0]
        if key not in LIB:
            continue
        py = LIB[key].iso()
        assert row[4] == pytest.approx(py["Y"], rel=1e-9), key
        assert row[8] == pytest.approx(py["kappa1"], rel=1e-9), key
        if LIB[key].back == "ext":
            assert row[3] == pytest.approx(py["U"], rel=1e-9), key
        seen += 1
    assert seen == len(LIB)


# --- Conduction Lab -----------------------------------------------------------------
def test_conduction_workbook_periodic_slab(soffice):
    """200 mm dense concrete, back insulated, room 24 +/- 4 C daily: after four
    days the peak surface flux is 4 |Y| with Y the ISO 13786 admittance of an
    insulated slab, (1+i)(lambda/delta) tanh((1+i) d/delta) behind Rsi."""
    def edit(wb):
        _custom(wb, [("dense concrete", 200)], "adiabatic")
        s = wb["Setup"]; s["B4"] = "Custom"; s["B10"] = 24
    ws = recalc("conduction-lab.xlsx", edit, soffice)["Results"]
    Y = Element("c", "c", "wall", [L("dense concrete", 0.2)], "adiabatic").iso()["Y"]
    assert ws["B6"].value == pytest.approx(4 * Y, rel=0.01)


def test_conduction_workbook_matches_python():
    """Default set-up (W2, 10 mm mesh): heat absorbed 09:00-18:00 on day 4
    equals the Python chain on the same mesh; Python's 2 mm reference mesh
    differs from it by under 0.1%."""
    from conduction import absorbed_kwh, hot_day_trace
    from elements import step_response
    ws = load_workbook(os.path.join(OUT, "conduction-lab.xlsx"), data_only=True)["Results"]
    t, Tr = hot_day_trace()
    h = (t / 3600) % 24
    sel = (t >= t[-1] - 86400 + 300) & (h >= 9) & (h < 18)
    for col, mode in (("B", "fd"), ("C", "5node")):
        q = step_response(LIB["W2"], Tr, 300.0, mode, dx=0.01, Tback=15.0)
        assert ws[f"{col}5"].value == pytest.approx(q[sel].sum() * 300 / 3600, rel=1e-6), mode
    assert ws["B5"].value == pytest.approx(absorbed_kwh(LIB["W2"], "fd") * 1000, rel=1e-3)


# --- House Simulator (5R1C) ----------------------------------------------------------
def test_house_workbook_matches_python(london):
    """Default inputs (dense block on dabs, hot year, no night vent)."""
    ws = load_workbook(os.path.join(OUT, "house-simulator.xlsx"), data_only=True)["Results"]
    rm = room.Room("Masonry (dense block, dabs), stud partitions")
    m = room.metrics(room5r1c.run(rm, london["hot"]))
    assert ws["B4"].value == pytest.approx(m["heat_kwh_m2"], rel=1e-6)
    assert ws["B8"].value == pytest.approx(m["dh26"], rel=1e-5)
    assert ws["B7"].value == m["tm59b_h"]


# --- Mass Sweeps ------------------------------------------------------------------------
def test_sweeps_workbook_interpolator():
    """The bilinear interpolator at t = 80 mm, 1.5 ACH equals a hand
    interpolation of sweep B between 50/100 mm and 1/2 ACH."""
    ws = load_workbook(os.path.join(OUT, "mass-sweeps.xlsx"), data_only=True)["Mass x night vent"]
    r = next(c.row for c in ws["A"] if c.value == "DH26 hot year")
    B = sweeps_json()["B"]; dh = np.array(B["dh26"])
    j0, j1 = B["t"].index(0.05), B["t"].index(0.1); i0, i1 = B["ach"].index(1), B["ach"].index(2)
    u, v = 0.6, 0.5
    expect = (1 - u) * (1 - v) * dh[i0, j0] + u * (1 - v) * dh[i0, j1] + (1 - u) * v * dh[i1, j0] + u * v * dh[i1, j1]
    assert ws.cell(r, 2).value == pytest.approx(expect, rel=1e-9)


def sweeps_json():
    import json
    return json.load(open(os.path.join(ROOT, "models", "results", "sweeps.json")))
